# SPDX-License-Identifier: Apache-2.0
"""Reproducible local-parcel exsolution bounds and source-data audit.

Run from the repository root. No physical clock or nucleation rate is fitted.
All compositions and lever-rule amounts use native (MgSiO3,H2) formula units.
"""

import hashlib
import itertools
import json
import platform
from datetime import datetime, timezone
from pathlib import Path
import re
import subprocess
import sys

import equinox as eqx
import jax
import jax.numpy as jnp
from molmass import Formula
import numpy as np
from openpyxl import load_workbook
import pandas as pd
import scipy
from scipy.optimize import brentq, minimize_scalar
from scipy.constants import Avogadro, elementary_charge, gas_constant

from atmodeller import ChemicalSpecies, EquilibriumModel, SolverParameters, SubregularSolution
from atmodeller.thermodata.defaults import default_properties
from research.spinodal_study.run_study import checks, portable

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
OUT = HERE / "outputs"
M_SIL = Formula("MgSiO3").mass
M_H2 = Formula("H2").mass


def save(name, value):
    (OUT / name).write_text(json.dumps(portable(value), indent=2, allow_nan=False) + "\n")


def to_x(w):
    """Native H2 mole fraction from local binary mass fraction."""
    w = np.asarray(w)
    return (w / M_H2) / (w / M_H2 + (1 - w) / M_SIL)


def to_w(x):
    x = np.asarray(x)
    return x * M_H2 / (x * M_H2 + (1 - x) * M_SIL)


def models():
    return {
        "gilmore26": SubregularSolution.hydrogen_silicate_gilmore26(),
        "rogers25": SubregularSolution.hydrogen_silicate_rogers25(),
        "wrong_basis_control": SubregularSolution(
            75000.0, 424000.0, temperature_scale=4400.0, pressure_scale=-350000.0
        ),
    }


@eqx.filter_jit
def onset_zero_pressure(model, x):
    """Invert the binodal; P scaling is analytical, so no P-grid root repeats.

    Zero pressure is an algebraic reference for this separable parameterization,
    not a claim that the published fit is physically calibrated at zero pressure.
    """
    lo = jnp.full_like(x, 250.0)
    hi = jnp.broadcast_to(model.temperature_scale, x.shape)

    def step(_, bounds):
        lower, upper = bounds
        mid = (lower + upper) / 2
        inside = model.coexistence(mid, 0.0).contains(x)
        return jnp.where(inside, mid, lower), jnp.where(inside, upper, mid)

    lo, hi = jax.lax.fori_loop(0, 38, step, (lo, hi))
    return (lo + hi) / 2


def critical(model):
    spinodal_t = eqx.filter_jit(model.spinodal_temperature)

    def objective(x):
        value = float(spinodal_t(0.0, x))
        return -value if np.isfinite(value) else 0.0

    r = minimize_scalar(
        objective, bounds=(0.001, 0.999), method="bounded", options={"xatol": 1e-13}
    )
    assert r.success and r.fun < 0
    return float(r.x), float(-r.fun)


def equilibrium_inventory(model, t, p, x):
    c = eqx.filter_jit(model.coexistence)(t, p)
    assert np.asarray(c.valid).all(), "Unresolved common tangent"
    fraction = np.asarray(c.upper_fraction(x))
    upper, lower = np.asarray(c.upper), np.asarray(c.lower)
    hydrogen = np.where(c.contains(x), fraction * upper / x, 0.0)
    reconstruction = (1 - fraction) * lower + fraction * upper
    error = np.where(c.contains(x), np.abs(reconstruction - x), 0.0)
    return hydrogen, fraction, lower, upper, error


def main_grid():
    weights = np.unique(
        np.r_[np.geomspace(0.001, 0.2, 161), [0.003, 0.01, 0.015, 0.02, 0.03, 0.05, 0.1]]
    )
    x = to_x(weights)
    rows, critical_rows, checks_rows = [], [], []
    pressure = np.linspace(2, 10, 41)
    for name, model in models().items():
        xc, tc0 = critical(model)
        onset = np.asarray(onset_zero_pressure(model, jnp.asarray(x)))
        ts0 = np.asarray(model.spinodal_temperature(0.0, x))
        # Reject a lost lower bracket rather than reporting the numerical bound as data.
        assert np.asarray(model.coexistence(jnp.full(x.shape, 250.0), 0.0).contains(x)).all()
        critical_rows.append(
            {
                "model": name,
                "critical_x_native": xc,
                "critical_z_fit": float(model.mixing_composition(xc)),
                "critical_hydrogen_wt_percent": 100 * float(to_w(xc)),
                "critical_temperature_zero_pressure_K": tc0,
            }
        )
        # Exact asymptotic lower composition for negative curvature while F>=0.
        a, b = float(model.interaction_0), float(model.interaction_1)
        zmin = (4 * a - 2 * b) / (6 * (a - b))
        critical_rows[-1]["minimum_spinodal_hydrogen_wt_percent_F_positive"] = 100 * float(
            to_w(model.native_composition(zmin))
        )
        for pg in pressure:
            p = pg * 1e4
            factor = 1 + p / float(model.pressure_scale)
            tb, ts = onset * factor, ts0 * factor
            # At onset, use a finite root only; otherwise this inventory is undefined.
            # Very low formal spinodals are outside the fluid study window and
            # can require exponentially unresolved phase endpoints. Do not
            # present extrapolated lever-rule inventories as physical results.
            resolved_onset = np.isfinite(ts) & (ts >= 2000.0)
            t_eval = np.where(resolved_onset, ts, 2500.0)
            h, frac, lower, upper, error = equilibrium_inventory(model, t_eval, p, x)
            for i, w in enumerate(weights):
                row = dict(
                    model=name,
                    pressure_GPa=pg,
                    hydrogen_wt_percent=100 * w,
                    native_x=x[i],
                    fit_z=float(model.mixing_composition(x[i])),
                    silicate_rich_bulk=bool(x[i] < xc),
                    binodal_temperature_K=tb[i],
                    spinodal_temperature_K=ts[i],
                    metastable_interval_K=tb[i] - ts[i],
                    H_in_rich_phase_at_spinodal=h[i] if resolved_onset[i] else np.nan,
                    spinodal_below_2500_K=bool(not np.isfinite(ts[i]) or ts[i] < 2500),
                    lever_rule_error=error[i],
                )
                rows.append(row)
            # Verify roots independently, including both sides of the critical composition.
            if pg in (2.0, 5.0, 10.0):
                for w in (0.01, 0.03, 0.05, 0.1):
                    i = int(np.argmin(abs(weights - w)))
                    t = tb[i]
                    c = model.coexistence(t, p)
                    endpoint = float(c.lower if x[i] < xc else c.upper)
                    row = dict(
                        model=name,
                        pressure_GPa=pg,
                        hydrogen_wt_percent=w * 100,
                        binodal_composition_error=abs(endpoint - x[i]),
                        above_binodal_two_phase=bool(
                            model.coexistence(t + 0.01, p).contains(x[i])
                        ),
                        below_binodal_two_phase=bool(
                            model.coexistence(t - 0.01, p).contains(x[i])
                        ),
                    )
                    if np.isfinite(ts[i]):
                        root = brentq(
                            lambda temperature: float(model.curvature(temperature, p, x[i])),
                            max(1.0, ts[i] - 0.5),
                            ts[i] + 0.5,
                            xtol=1e-8,
                        )
                        row["spinodal_temperature_independent_error_K"] = abs(root - ts[i])
                    checks_rows.append(row)
        print("completed onset map", name, flush=True)
    pd.DataFrame(rows).to_csv(OUT / "onsets.csv", index=False)
    pd.DataFrame(critical_rows).to_csv(OUT / "critical_points.csv", index=False)
    save("root_crosschecks.json", checks_rows)
    assert max(r["binodal_composition_error"] for r in checks_rows) < 1e-7
    assert all(
        not r["above_binodal_two_phase"] and r["below_binodal_two_phase"] for r in checks_rows
    )
    return pd.DataFrame(rows), critical_rows


def uncertainty_value(value):
    if isinstance(value, (float, int)):
        return float(value), 0.0
    match = re.fullmatch(r"(-?\d+(?:\.\d+)?)\((\d+)\)", value)
    assert match, value
    decimals = len(match[1].split(".")[1]) if "." in match[1] else 0
    return float(match[1]), float(match[2]) * 10**-decimals


def source_data_audit():
    ws = load_workbook(HERE / "sources/EDTable1.xlsx", data_only=True).active
    rows = []
    data_rows = [row for row in list(ws.values)[3:] if isinstance(row[1], (int, float))]
    assert len(data_rows) == 14, "Unexpected source table; inspect before interpreting"
    for data in data_rows:
        p, psigma = uncertainty_value(data[0])
        t = float(data[1])
        two = isinstance(data[5], str)
        if two:
            rich, poor = float(data[5].split("/")[-1]), float(data[6].split("/")[-1])
            rounded_atom_errors = []
            for field in (data[5], data[6]):
                atoms = [float(a) for a in field.split("/")]
                implied = (atoms[0] / 2) / (atoms[0] / 2 + 2 * sum(atoms[1:4]) / 5)
                # Atom counts are rounded and have a few-percent uncertainty.
                # Keep the author's reported composition as the measurement;
                # record this discrepancy instead of silently replacing it.
                rounded_atom_errors.append(abs(implied - atoms[-1]))
        else:
            rich = poor = abs(float(data[5]))
            rounded_atom_errors = [np.nan]
        for name, m in models().items():
            c = m.coexistence(t, p * 1e4)
            # Compare physical atom counts, not the similarly named but different x variables.
            # All external data are converted to native y=2z/(1+z).
            raw = np.array([poor, rich])
            native = 2 * raw / (1 + raw)
            predicted_native = np.array([float(c.lower), float(c.upper)])
            predicted_z = predicted_native / (2 - predicted_native)
            if two:
                diff = predicted_z - raw
                matches = bool(c.has_gap and np.max(np.abs(diff)) <= 0.05)
                # A common pressure must match BOTH endpoints. This is a
                # deterministic envelope over quoted P errors, not a likelihood.
                cp = eqx.filter_jit(m.coexistence)(
                    t, np.linspace(p - psigma, p + psigma, 101) * 1e4
                )
                yn = np.stack((np.asarray(cp.lower), np.asarray(cp.upper)), axis=-1)
                zp = yn / (2 - yn)
                residuals = np.max(np.abs(zp - raw), axis=-1)
                best_pressure_error = float(
                    np.min(np.where(cp.has_gap & cp.valid, residuals, np.inf))
                )
            else:
                diff = np.array([np.nan, np.nan])
                matches = not bool(c.contains(native[0]))
                best_pressure_error = np.nan
            rows.append(
                dict(
                    model=name,
                    pressure_GPa=p,
                    pressure_uncertainty_GPa=psigma,
                    temperature_K=t,
                    observed_two_phase=two,
                    observed_z_low=poor,
                    observed_z_high=rich,
                    predicted_z_low=predicted_z[0],
                    predicted_z_high=predicted_z[1],
                    predicted_has_gap=bool(c.has_gap),
                    endpoint_error_low=diff[0],
                    endpoint_error_high=diff[1],
                    maximum_composition_discrepancy_from_rounded_atoms=max(rounded_atom_errors),
                    matched_nominal_criterion=matches,
                    minimum_endpoint_max_error_over_pressure_uncertainty=best_pressure_error,
                    valid=bool(c.valid),
                )
            )
    df = pd.DataFrame(rows)
    df.to_csv(OUT / "source_data_audit.csv", index=False)
    return df


def homogeneous_data_audit():
    """Reconstruct excess properties from ED Table 2 on its oxide mixing basis.

    Pure endpoints are at 5.00 GPa, mixtures at 5.06-5.08 GPa. No equation
    of state correction is available; retain this small pressure offset.
    Uncertainties assume independent table entries; shared endmembers induce
    cross-composition covariance, so these residuals are not a chi-square fit.
    """
    ws = load_workbook(HERE / "sources/EDTable2.xlsx", data_only=True).active
    raw = [r for r in list(ws.values)[3:] if isinstance(r[1], (int, float))]
    assert len(raw) == 5
    rows = []
    for r in raw[1:-1]:
        nh, ns = (float(v) for v in r[2].split("/"))
        n_atoms, n_fit = nh + 5 * ns, nh / 2 + 2 * ns
        z = (nh / 2) / n_fit
        p, _ = uncertainty_value(r[0])
        h0 = z * (1 - z) * (75000 * (1 - z) + 424000 * z)
        row = dict(pressure_GPa=p, temperature_K=float(r[1]), fit_z=z)
        for field, index, unit, prediction in (
            ("enthalpy_J_per_fit_mol", 3, Avogadro * elementary_charge, h0 * (1 - p / 35)),
            ("excess_entropy_J_per_fit_mol_K", 4, gas_constant, h0 / 4400),
            ("volume_cm3_per_fit_mol", 5, Avogadro * 1e-24, -h0 / 35000),
        ):
            value, sigma = uncertainty_value(r[index])
            h, hs = uncertainty_value(raw[0][index])
            s, ss = uncertainty_value(raw[-1][index])
            excess = (value * n_atoms - nh * h - 5 * ns * s) / n_fit * unit
            uncertainty = (
                np.sqrt((sigma * n_atoms) ** 2 + (hs * nh) ** 2 + (ss * 5 * ns) ** 2)
                / n_fit
                * unit
            )
            if index == 4:
                excess += gas_constant * (z * np.log(z) + (1 - z) * np.log1p(-z))
            row.update(
                {
                    f"observed_{field}": excess,
                    f"sigma_{field}": uncertainty,
                    f"model_{field}": prediction,
                }
            )
        rows.append(row)
    pd.DataFrame(rows).to_csv(OUT / "homogeneous_data_audit.csv", index=False)


def stopping_temperatures():
    """Sensitivity to ending the fluid path before a formal instability occurs."""
    w = np.geomspace(0.001, 0.09, 301)
    m = models()["gilmore26"]
    rows = []
    for p, t in itertools.product((2.0, 5.0, 10.0), (2000.0, 2500.0, 3000.0)):
        h, _, _, _, _ = equilibrium_inventory(m, t, p * 1e4, to_x(w))
        ts = np.asarray(m.spinodal_temperature(p * 1e4, to_x(w)))
        for i, mass in enumerate(w):
            rows.append(
                dict(
                    pressure_GPa=p,
                    stopping_temperature_K=t,
                    hydrogen_wt_percent=100 * mass,
                    equilibrium_H_rich_fraction=h[i],
                    spinodal_reached=bool(np.isfinite(ts[i]) and ts[i] >= t),
                    withheld_H_fraction=0.0 if np.isfinite(ts[i]) and ts[i] >= t else h[i],
                )
            )
    pd.DataFrame(rows).to_csv(OUT / "stopping_temperatures.csv", index=False)


def cooling_paths():
    rows = []
    t = np.linspace(4000.0, 2000.0, 401)
    for name, m in models().items():
        for p, w in itertools.product((2.0, 5.0, 10.0), (0.01, 0.03, 0.05)):
            x = to_x(w)
            h, f, low, high, error = equilibrium_inventory(m, t, p * 1e4, x)
            ts = float(m.spinodal_temperature(p * 1e4, x))
            triggered = np.isfinite(ts) & (t < ts)
            delayed = np.where(triggered, h, 0.0)
            g = np.asarray(
                m.gibbs_energy(t, p * 1e4, x)
                - eqx.filter_jit(m.equilibrium_gibbs_energy)(t, p * 1e4, x)
            )
            assert min(g) >= -1e-7
            for i in range(len(t)):
                rows.append(
                    dict(
                        model=name,
                        pressure_GPa=p,
                        hydrogen_wt_percent=w * 100,
                        temperature_K=t[i],
                        H_equilibrium_fraction=h[i],
                        H_delayed_limit_fraction=delayed[i],
                        withheld_H_fraction=h[i] - delayed[i],
                        spinodal_reached=bool(triggered[i]),
                        equilibrium_energy_lowering_J_per_native_mol=g[i],
                        lower_native_x=low[i],
                        upper_native_x=high[i],
                        rich_phase_native_mole_fraction=f[i],
                        lever_rule_error=error[i],
                    )
                )
    df = pd.DataFrame(rows)
    df.to_csv(OUT / "cooling_paths.csv", index=False)
    return df


def sensitivity():
    factors = [(1.0, 1.0, 1.0, 1.0)]
    labels = ["nominal"]
    for i, key in enumerate(("A", "B", "C", "abs_D")):
        for f in (0.9, 1.1):
            v = [1.0] * 4
            v[i] = f
            factors.append(tuple(v))
            labels.append(f"{key}_{f}")
    for i, v in enumerate(itertools.product((0.9, 1.1), repeat=4)):
        factors.append(v)
        labels.append(f"corner_{i}")
    w = np.array([0.01, 0.015, 0.02, 0.03, 0.05])
    x = jnp.array(to_x(w))
    rows = []
    for label, factor in zip(labels, factors, strict=True):
        a, b, c, d = np.array([75000, 424000, 4400, -350000]) * factor
        m = SubregularSolution(
            a, b, temperature_scale=c, pressure_scale=d, component_0_multiplicity=2.0
        )
        tb0 = np.asarray(onset_zero_pressure(m, x))
        ts0 = np.asarray(m.spinodal_temperature(0.0, x))
        for p in (2.0, 5.0, 10.0):
            f = 1 + p * 1e4 / d
            tb, ts = tb0 * f, ts0 * f
            resolved_onset = np.isfinite(ts) & (ts >= 2000.0)
            h, *_ = equilibrium_inventory(
                m, np.where(resolved_onset, ts, 2500.0), p * 1e4, np.asarray(x)
            )
            for i, weight in enumerate(w):
                rows.append(
                    dict(
                        perturbation=label,
                        A_J_per_fit_mol=a,
                        B_J_per_fit_mol=b,
                        C_K=c,
                        D_bar=d,
                        pressure_GPa=p,
                        hydrogen_wt_percent=weight * 100,
                        binodal_temperature_K=tb[i],
                        spinodal_temperature_K=ts[i],
                        metastable_interval_K=tb[i] - ts[i],
                        H_in_rich_phase_at_spinodal=h[i] if resolved_onset[i] else np.nan,
                    )
                )
    df = pd.DataFrame(rows)
    df.to_csv(OUT / "sensitivity.csv", index=False)
    return df


def native_solver_check():
    rows = list(
        itertools.product(
            (2500.0, 3000.0, 3500.0), (20000.0, 50000.0, 100000.0), (0.01, 0.03, 0.05)
        )
    )
    t, p, w = np.array(rows).T
    species = (
        ChemicalSpecies.create_solution("MgSiO3"),
        ChemicalSpecies.create_solution("H2", state="fluid", thermo=default_properties["H2_g"]),
    )
    # Pure H2 standard state is irrelevant in a conserved closed binary. It is
    # NOT calibrated here for exchange with a separate atmospheric gas phase.
    m = SubregularSolution.hydrogen_silicate_gilmore26()
    model = EquilibriumModel.from_species(
        [],
        silicate_melt_species=species,
        temperature=t,
        pressure=p,
        mass=0.0,
        mass_constraints={"H": w, "MgSiO3": 1 - w},
        mixing_models={"silicate_melt": m},
        phase_equilibrium="spinodal",
        solver_parameters=SolverParameters(atol=1e-10, rtol=1e-10, max_steps=256),
    )
    output = model.solve_with_default()
    data = output.to_dict(to_numpy=True)
    validation = checks(output, data)
    assert np.asarray(validation["accepted"]).all()
    d = data["silicate_melt"]["phase"]["mixing"]
    np.testing.assert_allclose(
        np.asarray(d["bulk_component_1_mole_fraction"]).ravel(), to_x(w), atol=1e-10
    )
    expected = np.asarray(m.diagnostics(t, p, to_x(w))["equilibrium_two_phase"])
    np.testing.assert_array_equal(np.asarray(d["equilibrium_two_phase"]).ravel(), expected)
    save("native_solver_check.json", dict(cases=rows, validation=validation, mixing=d))
    np.save(OUT / "native_solver_solution.npy", np.asarray(output.solution))
    return {
        "number_solves": len(rows),
        "all_accepted": True,
        "maximum_relative_element_error": float(max(validation["maximum_relative_element_error"])),
    }


def main():
    OUT.mkdir(exist_ok=True)
    grid, crit = main_grid()
    audit = source_data_audit()
    homogeneous_data_audit()
    stopping_temperatures()
    paths = cooling_paths()
    stress = sensitivity()
    native = native_solver_check()
    summary = {
        "scope": "Closed local binary parcels; thermodynamic and no-nucleation limiting cases, not planetary ages or escape rates.",
        "critical_points": crit,
        "onset_rows": len(grid),
        "cooling_rows": len(paths),
        "stress_rows": len(stress),
        "native_solver": native,
        "maximum_lever_rule_error": float(paths.lever_rule_error.max()),
        "source_data_audit": audit.groupby("model")
        .apply(
            lambda d: {
                "two_phase_count": int(d.observed_two_phase.sum()),
                "two_phase_within_nominal_0_05": int(
                    d.loc[d.observed_two_phase, "matched_nominal_criterion"].sum()
                ),
                "two_phase_within_0_05_allowing_pressure_uncertainty": int(
                    (
                        d.loc[
                            d.observed_two_phase,
                            "minimum_endpoint_max_error_over_pressure_uncertainty",
                        ]
                        <= 0.05
                    ).sum()
                ),
                "single_phase_count": int((~d.observed_two_phase).sum()),
                "single_phase_agreements": int(
                    d.loc[~d.observed_two_phase, "matched_nominal_criterion"].sum()
                ),
                "maximum_endpoint_error": float(
                    d[["endpoint_error_low", "endpoint_error_high"]].abs().max().max()
                ),
            },
            include_groups=False,
        )
        .to_dict(),
        "examples": grid.loc[
            (grid.model == "gilmore26")
            & grid.pressure_GPa.isin([2.0, 5.0, 10.0])
            & np.any(
                np.isclose(grid.hydrogen_wt_percent.to_numpy()[:, None], [1.0, 3.0, 5.0]), axis=1
            )
        ].to_dict("records"),
    }
    save("summary.json", summary)
    sources = [
        *sorted((ROOT / "atmodeller").rglob("*.py")),
        Path(__file__),
        HERE / "PROTOCOL.md",
        ROOT / "uv.lock",
        ROOT / "research/spinodal_study/run_study.py",
        *sorted((HERE / "sources").glob("*")),
    ]
    save(
        "provenance.json",
        {
            "run_utc": datetime.now(timezone.utc).isoformat(),
            "git_head_before_work": subprocess.check_output(
                ["git", "rev-parse", "HEAD"], cwd=ROOT, text=True
            ).strip(),
            "native_molar_masses_g_per_mol": {"H2": M_H2, "MgSiO3": M_SIL},
            "sources": {
                str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest()
                for p in sources
            },
            "versions": {
                "python": sys.version,
                "platform": platform.platform(),
                "equinox": eqx.__version__,
                "scipy": scipy.__version__,
                "jax": jax.__version__,
                "numpy": np.__version__,
                "pandas": pd.__version__,
            },
        },
    )
    print(json.dumps(portable(summary), indent=2), flush=True)


if __name__ == "__main__":
    main()

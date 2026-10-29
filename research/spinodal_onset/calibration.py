# SPDX-License-Identifier: Apache-2.0
"""Recalibrate the H2-MgSiO3 excess free energy and propagate data uncertainty.

Post-audit analysis (protocol amendment, 29 October 2026). It does three things:

1. Quantifies how well the published four-parameter form can fit the source
   data jointly: Extended Data Table 1 coexistence and single-phase runs, and
   Table 2 homogeneous excess enthalpy, entropy and volume with their full
   covariance (shared endmembers and the 5.00 vs 5.06-5.08 GPa offset).
2. Replaces the +/-10% parameter stress test by posterior distributions of the
   delayed-exsolution quantities, for four smooth (Redlich-Kister) model classes.
3. Records how strongly those quantities depend on the model class, i.e. how
   far the existing data constrain the metastable free-energy curvature.

The likelihood has three parts. (a) Phase state: each of the 14 coexistence
runs is one or two phases at its (T, P). A misclassified run costs
((T_run - T_b)/sigma_T)^2, with T_b the model binodal temperature at the run's
bulk composition and sigma_T^2 = (50 K)^2 + (100 K/GPa * sigma_P)^2; this is
the most robust information in the source. (b) Coexisting compositions, quoted
+/-0.05 treated as 1-sigma Gaussian. (c) Homogeneous excess properties with
full covariance. Priors are uniform within the stated bounds, with inverse tau
and pi sampled. Run from the repository root; ~20 min on 10 cores.
"""

from concurrent.futures import ProcessPoolExecutor
from dataclasses import dataclass
import json
import math
from pathlib import Path
import re
import sys
import time

import numpy as np
from openpyxl import load_workbook
import pandas as pd
from scipy.constants import Avogadro, elementary_charge, gas_constant
from scipy.optimize import minimize

from research.spinodal_onset.mixing_models import ExcessModel, weight_to_z

HERE = Path(__file__).resolve().parent
OUT = HERE / "outputs"
SIGMA_X = 0.05
EV = Avogadro * elementary_charge  # J/mol per eV/particle
A3 = Avogadro * 1e-24  # cm3/mol per A^3/particle
TARGET = dict(weight=0.03, pressure_GPa=2.0)  # headline case of the study
SEED = 20261029


def uncertainty_value(value):
    """Parse '2.05(4)' -> (2.05, 0.04); plain numbers have zero uncertainty."""
    if isinstance(value, (float, int)):
        return float(value), 0.0
    match = re.fullmatch(r"(-?\d+(?:\.\d+)?)\((\d+)\)", value)
    assert match, value
    decimals = len(match[1].split(".")[1]) if "." in match[1] else 0
    return float(match[1]), float(match[2]) * 10**-decimals


# ----------------------------------------------------------------------------- data


def coexistence_runs():
    ws = load_workbook(HERE / "sources/EDTable1.xlsx", data_only=True).active
    rows = [r for r in list(ws.values)[3:] if isinstance(r[1], (int, float))]
    assert len(rows) == 14
    two, one = [], []
    for r in rows:
        p, sp = uncertainty_value(r[0])
        nh = float(r[2].split("/")[0])
        bulk = (nh / 2) / (nh / 2 + 54)  # 27 MgSiO3 = 54 oxide units
        if isinstance(r[5], str):
            hi, lo = float(r[5].split("/")[-1]), float(r[6].split("/")[-1])
            two.append(dict(p=p, sp=sp, t=float(r[1]), lo=lo, hi=hi, bulk=bulk))
        else:
            one.append(dict(p=p, sp=sp, t=float(r[1]), x=abs(float(r[5])), bulk=bulk))
    return two, one


def figure_values():
    """Digitized Fig. 2a compositions for the two runs where table and figure differ."""
    d = json.loads((OUT / "source_figure_digitized.json").read_text())
    out = {}
    for series, t in (("red", 3500), ("grey", 3000)):
        xs = sorted(
            m["x"]
            for m in d["markers"]
            if m["series"] == series and abs(m["temperature_K"] - t) < 20
        )
        out[t if series == "red" else -t] = xs
    return out  # keys: 3500 -> 1.65 GPa run, -3000 -> 5.14 GPa run


def homogeneous_data():
    """Excess H (J/mol), S (J/mol/K), V (cm3/mol) per fit mole with covariance.

    Mixture entries are independent; endmember errors are shared, giving
    off-diagonal covariance. Endmembers were run at 5.00 GPa and mixtures at
    5.06-5.08 GPa: a fully correlated systematic of V*dP (enthalpy) and
    V*dP/K (volume, K=15 GPa for H fluid, 30 GPa for silicate) is added.
    """
    ws = load_workbook(HERE / "sources/EDTable2.xlsx", data_only=True).active
    raw = [r for r in list(ws.values)[3:] if isinstance(r[1], (int, float))]
    assert len(raw) == 5 and raw[0][2] == "160/0" and raw[-1][2] == "0/27"
    pure_h, pure_s, mixes = raw[0], raw[-1], raw[1:-1]
    vh, _ = uncertainty_value(pure_h[5])
    vs, _ = uncertainty_value(pure_s[5])
    states, values, blocks, systematic = [], [], [], []
    for index, unit, label in ((3, EV, "H"), (4, gas_constant, "S"), (5, A3, "V")):
        mean, jac_mix, jac_end = [], [], []
        sh = uncertainty_value(pure_h[index])
        ss = uncertainty_value(pure_s[index])
        shift = []
        for r in mixes:
            nh, ns = (float(v) for v in r[2].split("/"))
            n_atoms, n_fit = nh + 5 * ns, nh / 2 + 2 * ns
            z = (nh / 2) / n_fit
            p, _ = uncertainty_value(r[0])
            value, sigma = uncertainty_value(r[index])
            e = (value * n_atoms - nh * sh[0] - 5 * ns * ss[0]) / n_fit * unit
            if label == "S":  # total -> excess: remove ideal mixing entropy
                e += gas_constant * (z * math.log(z) + (1 - z) * math.log1p(-z))
            mean.append(e)
            jac_mix.append(sigma * n_atoms / n_fit * unit)
            jac_end.append((nh * sh[1] / n_fit * unit, 5 * ns * ss[1] / n_fit * unit))
            dp = p - 5.0
            if label == "H":  # eV: A^3*GPa = 1e-21 J
                shift.append(
                    (nh * vh + 5 * ns * vs) * dp * 1e-21 / elementary_charge / n_fit * unit
                )
            elif label == "V":
                shift.append((nh * vh / 15 + 5 * ns * vs / 30) * dp / n_fit * unit)
            else:
                shift.append(0.0)
            if label == "H":
                states.append(dict(z=z, p=p, t=float(r[1])))
        cov = np.diag(np.square(jac_mix))
        a = np.array([j[0] for j in jac_end])
        b = np.array([j[1] for j in jac_end])
        cov += np.outer(a, a) + np.outer(b, b) + np.outer(shift, shift)
        values += mean
        blocks.append(cov)
        systematic.append(shift)
    covariance = np.zeros((9, 9))
    for i, block in enumerate(blocks):
        covariance[3 * i : 3 * i + 3, 3 * i : 3 * i + 3] = block
    return states, np.array(values), covariance


# ----------------------------------------------------------------------------- models


@dataclass(frozen=True)
class ModelClass:
    name: str
    labels: tuple
    lower: tuple
    upper: tuple
    start: tuple

    def build(self, theta):
        if self.name == "M1":
            a, b, it, ip = theta
            return ExcessModel.published_form(a * 1e3, b * 1e3, _inv(it), _inv(ip))
        if self.name == "M2":
            ah, bh, as_, bs, av, bv = theta
            return ExcessModel.subregular_hsv(ah * 1e3, bh * 1e3, as_, bs, av * 1e3, bv * 1e3)
        if self.name in ("M3", "M4"):
            *ls, it, ip = theta
            l0, l1, *extra = (v * 1e3 for v in ls)
            # published_form takes (A, B); invert L0=(A+B)/2, L1=(A-B)/2.
            return ExcessModel.published_form(
                l0 + l1, l0 - l1, _inv(it), _inv(ip), extra=tuple(extra)
            )
        raise ValueError(self.name)


def _inv(value):
    return math.inf if value == 0 else 1 / value


# Bounds: broad and stated. it = 1/tau (1/K), ip = 1/pi (1/GPa).
CLASSES = {
    "M1": ModelClass(
        "M1",
        ("A_kJ", "B_kJ", "inv_tau_per_K", "inv_pi_per_GPa"),
        (-300, -300, 0.0, -0.25),
        (1500, 3000, 1e-3, 0.1),
        (75, 424, 1 / 4400, -1 / 35),
    ),
    "M2": ModelClass(
        "M2",
        ("AH_kJ", "BH_kJ", "AS_J_per_K", "BS_J_per_K", "AV_cm3", "BV_cm3"),
        (-300, -300, -200, -200, -60, -60),
        (3000, 3000, 800, 800, 30, 30),
        (75, 424, 75 / 4.4, 424 / 4.4, -75 / 35, -424 / 35),
    ),
    "M3": ModelClass(
        "M3",
        ("L0_kJ", "L1_kJ", "L2_kJ", "inv_tau_per_K", "inv_pi_per_GPa"),
        (-300, -2000, -2000, 0.0, -0.25),
        (3000, 2000, 2000, 1e-3, 0.1),
        (249.5, -174.5, 0.0, 1 / 4400, -1 / 35),
    ),
    "M4": ModelClass(
        "M4",
        ("L0_kJ", "L1_kJ", "L2_kJ", "L3_kJ", "inv_tau_per_K", "inv_pi_per_GPa"),
        (-300, -2000, -2000, -2000, 0.0, -0.25),
        (3000, 2000, 2000, 2000, 1e-3, 0.1),
        (249.5, -174.5, 0.0, 0.0, 1 / 4400, -1 / 35),
    ),
}
VERSIONS = {  # data treatment variants; every version keeps the 14 phase states
    "table": "all data; ED Table 1 compositions, sigma_x = 0.05",
    "figure": "all data; Fig. 2a compositions for the two near-critical runs",
    "inflated": "all data; sigma_x = 0.15 for the two near-critical runs",
    "no_volume": "all data except the homogeneous excess volumes",
    "phase_comp": "phase states and compositions only (no homogeneous data)",
    "phase_hom": "phase states and homogeneous data only (no compositions)",
}
RUNS = [("M1", v) for v in VERSIONS] + [("M2", "table"), ("M3", "table"), ("M4", "table")]
SIGMA_T = 50.0  # K: finite-size/finite-time classification tolerance
DTC_DP = 100.0  # K/GPa: converts quoted pressure errors into temperature


class Problem:
    def __init__(self, model, version):
        self.cls = CLASSES[model]
        self.version = version
        self.two, self.one = coexistence_runs()
        self.states, self.hom, cov = homogeneous_data()
        self.keep = {"no_volume": np.arange(6), "phase_comp": np.arange(0)}.get(
            version, np.arange(9)
        )
        self.use_compositions = version != "phase_hom"
        self.hom_inv = np.linalg.inv(cov[np.ix_(self.keep, self.keep)]) if len(self.keep) else None
        fig = figure_values()
        self.sigma = []
        for run in self.two:
            near_critical = run["t"] == 3500 or (run["t"] == 3000 and run["p"] > 4)
            run["near_critical"] = near_critical
            if version == "figure" and near_critical:
                run["lo"], run["hi"] = fig[3500 if run["t"] == 3500 else -3000]
            self.sigma.append(0.15 if version == "inflated" and near_critical else SIGMA_X)
        self.dim = len(self.cls.labels)

    def start(self):
        return np.array(self.cls.start, float)

    @staticmethod
    def predict_two_phase(model, t, p):
        gap = model.binodal(t, p)
        if gap is None:  # one phase: both endpoints at the least stable composition
            zs = np.linspace(0.01, 0.99, 197)
            z = float(zs[np.argmin([model.d2f(v, t, p) for v in zs])])
            return z, z, False
        return gap[0], gap[1], True

    @staticmethod
    def misclassification_K(model, run, two_phase):
        """Temperature by which the model misplaces a run's phase state (0 if right)."""
        gap = model.binodal(run["t"], run["p"])
        inside = gap is not None and gap[0] < run["bulk"] < gap[1]
        if inside == two_phase:
            return 0.0
        tb = model.binodal_temperature(run["bulk"], run["p"])
        if not math.isfinite(tb):  # never two-phase at this composition
            return run["t"] if two_phase else 0.0
        return abs(run["t"] - tb)

    def chi2_parts(self, theta):
        model = self.cls.build(theta)
        parts = dict(phase_state=0.0, coexistence=0.0, homogeneous=0.0)
        for two_phase, runs in ((True, self.two), (False, self.one)):
            for run in runs:
                sigma = math.hypot(SIGMA_T, DTC_DP * run["sp"])
                parts["phase_state"] += (
                    self.misclassification_K(model, run, two_phase) / sigma
                ) ** 2
        for run, s in zip(self.two, self.sigma):
            if not self.use_compositions:
                break
            lo, hi, _ = self.predict_two_phase(model, run["t"], run["p"])
            parts["coexistence"] += ((lo - run["lo"]) / s) ** 2 + ((hi - run["hi"]) / s) ** 2
        if self.hom_inv is None:
            return parts
        pred = [model.enthalpy(st["z"], st["p"]) for st in self.states]
        pred += [model.entropy(st["z"]) for st in self.states]
        pred += [model.volume_cm3(st["z"]) for st in self.states]
        r = (self.hom - np.array(pred))[self.keep]
        parts["homogeneous"] = float(r @ self.hom_inv @ r)
        return parts

    def log_post(self, theta):
        theta = np.asarray(theta, float)
        lo, hi = self.cls.lower, self.cls.upper
        if any(not (a <= v <= b) for v, a, b in zip(theta, lo, hi)):
            return -math.inf
        try:
            parts = self.chi2_parts(theta)
        except (ValueError, ZeroDivisionError, OverflowError, RuntimeError):
            return -math.inf  # e.g. three-phase topology: outside the model class
        total = sum(parts.values())
        return -0.5 * total if math.isfinite(total) else -math.inf


# ----------------------------------------------------------------------------- derived


def derived(model):
    """Delayed-exsolution quantities at the target case, or NaN if undefined."""
    z0, p = weight_to_z(TARGET["weight"]), TARGET["pressure_GPa"]
    out = dict(
        T_C_K=math.nan,
        z_C=math.nan,
        binodal_K=math.nan,
        spinodal_K=math.nan,
        interval_K=math.nan,
        H_rich_fraction_at_spinodal=math.nan,
        spinodal_threshold_wt=math.nan,
    )
    try:
        tc, zc = model.critical_point(p)
        out.update(T_C_K=tc, z_C=zc)
        ts = model.spinodal_temperature(z0, p)
        tb = model.binodal_temperature(z0, p)
        out.update(binodal_K=tb, spinodal_K=ts, interval_K=tb - ts)
        if math.isfinite(ts):
            out["H_rich_fraction_at_spinodal"] = model.rich_hydrogen_fraction(z0, ts, p)
        ws = np.linspace(0.001, 0.2, 400)
        finite = [w for w in ws if math.isfinite(model.spinodal_temperature(weight_to_z(w), p))]
        out["spinodal_threshold_wt"] = 100 * finite[0] if finite else math.nan
    except (ValueError, RuntimeError):
        pass
    return out


# ----------------------------------------------------------------------------- sampling


def _optimize(args):
    model, version, x = args
    problem = Problem(model, version)
    best = None
    for tol in ((1e-8, 1e-6), (1e-9, 1e-8), (1e-10, 1e-9)):
        r = minimize(
            lambda v: -problem.log_post(v),
            x if best is None else best.x,
            method="Nelder-Mead",
            options=dict(maxiter=4000, maxfev=4000, xatol=tol[0], fatol=tol[1], adaptive=True),
        )
        best = r if best is None or r.fun <= best.fun else best
    return best.x, best.fun


def find_map(problem, pool, starts=8, seed=SEED):
    """Best of several Nelder-Mead restarts around the published parameters."""
    rng = np.random.default_rng(seed)
    x0 = problem.start()
    scale = np.array([(b - a) * 0.02 for a, b in zip(problem.cls.lower, problem.cls.upper)])
    candidates = [x0]
    while len(candidates) < starts:
        x = x0 + rng.normal(0, 1, problem.dim) * scale
        if math.isfinite(problem.log_post(x)):
            candidates.append(x)
    results = list(
        pool.map(_optimize, [(problem.cls.name, problem.version, x) for x in candidates])
    )
    x, fun = min(results, key=lambda r: r[1])
    return x, -fun, scale


def chain(args):
    model, version, seed, x0, scale, n_burn, n_keep, thin = args
    problem = Problem(model, version)
    rng = np.random.default_rng(seed)
    d = problem.dim
    x = np.array(x0, float) + rng.normal(0, 0.1, d) * scale
    lp = problem.log_post(x)
    while not math.isfinite(lp):
        x = np.array(x0, float) + rng.normal(0, 0.05, d) * scale
        lp = problem.log_post(x)
    cov = np.diag(np.square(scale * 0.05))
    factor, history, accepted, kept = 2.38**2 / d, [], 0, []
    window_acc = 0
    for i in range(n_burn + n_keep):
        prop = x + np.linalg.cholesky(factor * cov + 1e-14 * np.eye(d)) @ rng.normal(size=d)
        lpp = problem.log_post(prop)
        ok = math.log(rng.random()) < lpp - lp
        if ok:
            x, lp = prop, lpp
            window_acc += 1
            if i >= n_burn:
                accepted += 1
        history.append(x.copy())
        if i < n_burn and i >= 400 and i % 200 == 0:  # adapt during burn-in only
            cov = np.cov(np.array(history[len(history) // 2 :]).T) + 1e-12 * np.diag(scale**2)
            factor *= math.exp(window_acc / 200 - 0.234)
            window_acc = 0
        if i >= n_burn and (i - n_burn) % thin == 0:
            kept.append((x.copy(), lp))
    rows = []
    for theta, lp in kept:
        m = problem.cls.build(theta)
        row = {k: float(v) for k, v in zip(problem.cls.labels, theta)}
        row.update(log_post=lp, **derived(m))
        rows.append(row)
    trace = np.array(history[n_burn:])
    return dict(rows=rows, acceptance=accepted / n_keep, trace=trace[::10].tolist())


def split_rhat(chains):
    """Split-chain Gelman-Rubin statistic for one scalar quantity."""
    halves = []
    for c in chains:
        n = len(c) // 2
        halves += [c[:n], c[n : 2 * n]]
    halves = np.array(halves)
    n = halves.shape[1]
    w = halves.var(axis=1, ddof=1).mean()
    b = n * halves.mean(axis=1).var(ddof=1)
    return float(math.sqrt(((n - 1) / n * w + b / n) / w)) if w > 0 else math.nan


def ess(chains):
    """Effective sample size from the initial positive autocorrelation sequence."""
    total = 0.0
    for c in chains:
        c = np.asarray(c) - np.mean(c)
        n = len(c)
        if np.var(c) == 0:
            continue
        ac = np.correlate(c, c, "full")[n - 1 :] / (np.var(c) * np.arange(n, 0, -1))
        s = 0.0
        for k in range(1, n // 2):
            if ac[k] + ac[k + 1] < 0:
                break
            s += ac[k]
        total += n / (1 + 2 * s)
    return float(total)


def run(model, version, pool, n_chains=8, n_burn=3000, n_keep=9000, thin=30):
    problem = Problem(model, version)
    t0 = time.time()
    x_map, lp_map, scale = find_map(problem, pool)
    parts = problem.chi2_parts(x_map)
    jobs = [
        (model, version, SEED + 97 * i, x_map, scale, n_burn, n_keep, thin)
        for i in range(n_chains)
    ]
    results = list(pool.map(chain, jobs))
    rows = []
    for i, r in enumerate(results):
        for row in r["rows"]:
            rows.append(dict(model=model, version=version, chain=i, **row))
    df = pd.DataFrame(rows)
    keys = list(problem.cls.labels) + ["interval_K", "spinodal_K", "binodal_K", "T_C_K"]
    diag = {}
    for k in keys:
        series = [df[df.chain == i][k].to_numpy() for i in range(n_chains)]
        series = [s[np.isfinite(s)] for s in series]
        diag[k] = dict(rhat=split_rhat(series), ess=ess(series))
    traces = [np.array(r["trace"]) for r in results]
    trace_rhat = {
        lab: split_rhat([t[:, j] for t in traces]) for j, lab in enumerate(problem.cls.labels)
    }
    m = problem.cls.build(x_map)
    pred = []
    for run_ in problem.two:
        lo, hi, gap = problem.predict_two_phase(m, run_["t"], run_["p"])
        pred.append(
            dict(
                t=run_["t"],
                p=run_["p"],
                obs_lo=run_["lo"],
                obs_hi=run_["hi"],
                pred_lo=lo,
                pred_hi=hi,
                has_gap=gap,
                misclassified_K=problem.misclassification_K(m, run_, True),
            )
        )
    single = [
        dict(
            t=r["t"],
            p=r["p"],
            bulk=r["bulk"],
            misclassified_K=problem.misclassification_K(m, r, False),
        )
        for r in problem.one
    ]
    summary = dict(
        model=model,
        version=version,
        description=VERSIONS[version],
        parameters=list(problem.cls.labels),
        map=dict(zip(problem.cls.labels, map(float, x_map))),
        map_log_post=float(lp_map),
        map_chi2=parts,
        map_chi2_total=float(sum(parts.values())),
        n_data=dict(
            phase_states=len(problem.two) + len(problem.one),
            coexistence=8 if problem.use_compositions else 0,
            homogeneous=len(problem.keep),
        ),
        map_derived=derived(m),
        map_two_phase=pred,
        map_single_phase=single,
        acceptance=[r["acceptance"] for r in results],
        diagnostics=diag,
        parameter_trace_rhat=trace_rhat,
        quantiles={
            k: _quantiles(df[k])
            for k in (
                "interval_K",
                "spinodal_K",
                "binodal_K",
                "H_rich_fraction_at_spinodal",
                "T_C_K",
                "z_C",
                "spinodal_threshold_wt",
                *problem.cls.labels,
            )
        },
        undefined_spinodal_fraction=float(np.mean(~np.isfinite(df["spinodal_K"]))),
        seconds=time.time() - t0,
    )
    print(
        f"{model}/{version}: chi2_MAP={summary['map_chi2_total']:.2f} "
        f"interval={summary['quantiles']['interval_K']} "
        f"acc={np.round(summary['acceptance'], 2)} t={summary['seconds']:.0f}s",
        flush=True,
    )
    return summary, df


def _quantiles(series):
    s = np.asarray(series, float)
    s = s[np.isfinite(s)]
    if not len(s):
        return None
    q = np.percentile(s, [2.5, 16, 50, 84, 97.5])
    return dict(zip(("p2.5", "p16", "p50", "p84", "p97.5"), map(float, q)), n=int(len(s)))


def main():
    selected = RUNS if len(sys.argv) == 1 else [tuple(a.split("/")) for a in sys.argv[1:]]
    summaries, frames = [], []
    with ProcessPoolExecutor(max_workers=8) as pool:
        for model, version in selected:
            s, df = run(model, version, pool)
            summaries.append(s)
            frames.append(df)
    published = derived(ExcessModel.published_form(75e3, 424e3, 4400.0, -35.0))
    p = Problem("M1", "table")
    published_chi2 = p.chi2_parts(np.array([75, 424, 1 / 4400, -1 / 35]))
    states, values, cov = homogeneous_data()
    result = dict(
        protocol="Post-audit recalibration; see PROTOCOL.md amendment of 29 October 2026.",
        target=TARGET,
        sigma_x=SIGMA_X,
        sigma_T_K=SIGMA_T,
        dTC_dP_K_per_GPa=DTC_DP,
        versions=VERSIONS,
        homogeneous_data=dict(
            states=states,
            values=values.tolist(),
            covariance=cov.tolist(),
            order="H(J/mol) x3, S_excess(J/mol/K) x3, V(cm3/mol) x3",
        ),
        published_fit=dict(
            derived=published, chi2=published_chi2, chi2_total=float(sum(published_chi2.values()))
        ),
        runs=summaries,
    )
    (OUT / "calibration_summary.json").write_text(
        json.dumps(result, indent=2, default=float) + "\n"
    )
    pd.concat(frames).to_csv(OUT / "calibration_samples.csv", index=False)


if __name__ == "__main__":
    main()
